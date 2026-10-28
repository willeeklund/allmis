from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

ALLMIS = re.compile(r"allmis|open[\s-]?ice", re.IGNORECASE)
TIME_RANGE = re.compile(r"(\d{1,2})[.:](\d{2})\s*[-–]\s*(\d{1,2})[.:](\d{2})")
WEEK = re.compile(r"vecka(?:20|[\s_-])*(\d{1,2})(?!\d)", re.IGNORECASE)
YEAR = re.compile(r"(?<!\d)(20\d{2})(?!\d)")
DAY_LETTERS = {"M", "T", "O", "F", "L", "S"}
SLOT_MINUTES = 30


class ParseError(Exception):
    pass


@dataclass(frozen=True)
class Session:
    date: dt.date
    start: str
    end: str
    hall: str
    note: str = ""

    def to_dict(self) -> dict:
        return {
            "date": self.date.isoformat(),
            "start": self.start,
            "end": self.end,
            "hall": self.hall,
            "note": self.note,
        }


@dataclass
class Schedule:
    sessions: list[Session]
    first_day: dt.date
    last_day: dt.date


def week_from_name(name: str) -> int:
    m = WEEK.search(name)
    if not m:
        raise ParseError(f"Hittar inget veckonummer i filnamnet {name!r}")
    return int(m.group(1))


def schedule_start(name: str, today: dt.date) -> dt.date:
    week = week_from_name(name)
    year_match = YEAR.search(WEEK.sub("", name))
    if year_match:
        return dt.date.fromisocalendar(int(year_match.group(1)), week, 1)
    candidates = []
    for year in (today.year - 1, today.year, today.year + 1):
        try:
            candidates.append(dt.date.fromisocalendar(year, week, 1))
        except ValueError:
            pass
    return min(candidates, key=lambda d: abs((d - today).days))


def _fill(cell) -> str | None:
    f = cell.fill
    if f is None or not f.fill_type:
        return None
    color = f.fgColor
    if color.type == "rgb":
        return color.rgb
    return f"{color.type}:{color.theme if color.type == 'theme' else color.indexed}:{color.tint}"


def _minutes_to_hhmm(minutes: int) -> str:
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def _hour_columns(ws) -> dict[int, int]:
    return {
        c.column: c.value
        for c in ws[1]
        if isinstance(c.value, int) and not isinstance(c.value, bool) and c.column >= 3
    }


def _is_grid_sheet(ws) -> bool:
    return isinstance(ws.cell(1, 1).value, str) and len(_hour_columns(ws)) >= 3


def _column_to_minutes(ws):
    hours = _hour_columns(ws)
    first_col = min(hours)
    first_hour = hours[first_col]
    for col, hour in hours.items():
        if (col - first_col) * SLOT_MINUTES != (hour - first_hour) * 60:
            raise ParseError(
                f"Fliken {ws.title!r}: oväntad tidsrubrik {hour} i kolumn {col}"
            )
    return lambda col: first_hour * 60 + (col - first_col) * SLOT_MINUTES


def _note(text: str) -> str:
    rest = TIME_RANGE.sub("", ALLMIS.sub("", text))
    return re.sub(r"\s+", " ", rest).strip(" ,.-–")


def _parse_sheet(ws, start: dt.date) -> tuple[list[Session], list[dt.date]]:
    hall = ws.cell(1, 1).value.strip()
    to_minutes = _column_to_minutes(ws)
    sessions: list[Session] = []
    days: list[dt.date] = []
    cursor = start

    for r in range(2, ws.max_row + 1):
        letter = ws.cell(r, 1).value
        day = ws.cell(r, 2).value
        if not (isinstance(letter, str) and letter.strip() in DAY_LETTERS and isinstance(day, int)):
            continue
        for _ in range(62):
            if cursor.day == day:
                break
            cursor += dt.timedelta(days=1)
        else:
            raise ParseError(f"Fliken {ws.title!r} rad {r}: hittar inget datum för dag {day}")
        days.append(cursor)

        for col in range(3, ws.max_column + 1):
            cell = ws.cell(r, col)
            if not (isinstance(cell.value, str) and ALLMIS.search(cell.value)):
                continue
            fill = _fill(cell)
            end_col = col
            while end_col < ws.max_column:
                nxt = ws.cell(r, end_col + 1)
                if nxt.value is not None or _fill(nxt) != fill:
                    break
                end_col += 1
            explicit = TIME_RANGE.search(cell.value)
            if explicit:
                h1, m1, h2, m2 = (int(g) for g in explicit.groups())
                begin, finish = h1 * 60 + m1, h2 * 60 + m2
            else:
                begin, finish = to_minutes(col), to_minutes(end_col + 1)
            sessions.append(
                Session(cursor, _minutes_to_hhmm(begin), _minutes_to_hhmm(finish), hall, _note(cell.value))
            )
    return sessions, days


def merge_sessions(sessions: list[Session]) -> list[Session]:
    merged: list[Session] = []
    for s in sorted(set(sessions), key=lambda s: (s.date, s.hall, s.start, s.end)):
        prev = merged[-1] if merged else None
        if prev and prev.date == s.date and prev.hall == s.hall and s.start <= prev.end:
            notes = [n for n in dict.fromkeys([prev.note, s.note]) if n]
            merged[-1] = Session(prev.date, prev.start, max(prev.end, s.end), prev.hall, ", ".join(notes))
        else:
            merged.append(s)
    return sorted(merged, key=lambda s: (s.date, s.start, s.hall))


def parse_workbook(path: Path | str, name: str | None = None, today: dt.date | None = None) -> Schedule:
    path = Path(path)
    name = name or path.name
    today = today or dt.date.today()
    start = schedule_start(name, today)

    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        raise ParseError(f"Kan inte öppna {name!r} som Excel: {e}") from e

    grids = [ws for ws in wb.worksheets if _is_grid_sheet(ws)]
    if not grids:
        raise ParseError(f"Hittar ingen flik med tidsrutnät i {name!r}")

    sessions: list[Session] = []
    days: list[dt.date] = []
    for ws in grids:
        s, d = _parse_sheet(ws, start)
        sessions += s
        days += d
    if not days:
        raise ParseError(f"Hittar inga dagar i {name!r}")

    return Schedule(merge_sessions(sessions), min(days), max(days))
